"""Извлечение содержимого из файлов. Каждая функция возвращает список
(loc, label, snippet, payload), где payload идёт в модель."""
import hashlib
import locale
import os
import re

import numpy as np
from PIL import Image, ImageOps

Image.MAX_IMAGE_PIXELS = 200_000_000
ANSI = locale.getencoding()  # кодовая страница ANSI системы (cp1251, cp1252…): так сохраняют старые редакторы


class Skip(Exception):
    """Файл осознанно пропущен (слишком большой, бинарный, пустой...)."""


def _chunks(text, size=1500, overlap=200):
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    if not text:
        return []
    out, step = [], size - overlap
    for i in range(0, max(len(text) - overlap, 1), step):
        piece = text[i:i + size].strip()
        if len(piece) > 20:
            out.append(piece)
    return out


def _doc_items(name, pieces_with_labels, max_chunks):
    items = []
    for loc, (label, piece) in enumerate(pieces_with_labels[:max_chunks]):
        items.append((loc, label, piece[:300], f"title: {name} | text: {piece}"))
    return items


def _read_plain(path, cfg):
    if os.path.getsize(path) > cfg["max_text_mb"] * 1024 * 1024:
        raise Skip("слишком большой")
    with open(path, "rb") as f:
        raw = f.read()
    if b"\x00" in raw[:4096]:
        raise Skip("бинарный")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode(ANSI, errors="replace")


def extract_text(path, cfg):
    name = os.path.basename(path)
    ext = os.path.splitext(path)[1].lower()
    labeled = []
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        for i, page in enumerate(reader.pages[:cfg["max_pdf_pages"]]):
            try:
                t = page.extract_text() or ""
            except Exception:
                continue
            labeled += [(f"стр. {i + 1}", p) for p in _chunks(t)]
    elif ext == ".docx":
        import docx

        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        labeled = [(f"фрагмент {i + 1}", p) for i, p in enumerate(_chunks("\n".join(parts)))]
    else:
        if ".min." in name:
            raise Skip("минифицированный")
        labeled = [(f"фрагмент {i + 1}", p) for i, p in enumerate(_chunks(_read_plain(path, cfg)))]
    if not labeled:
        raise Skip("нет текста")
    return _doc_items(name, labeled, cfg["max_chunks_per_file"])


def load_image(path, max_side=768):
    im = Image.open(path)
    if im.format == "JPEG":
        im.draft("RGB", (max_side, max_side))  # декодирует сразу в уменьшенном размере — в разы быстрее
    im = ImageOps.exif_transpose(im)
    im = im.convert("RGB")
    im.thumbnail((max_side, max_side))
    return im


def extract_image(path, cfg):
    with Image.open(path) as probe:
        if min(probe.size) < cfg["min_image_side"]:
            raise Skip("маленькая картинка")
    return [(0, "", "", load_image(path))]


def _ts(sec):
    sec = int(sec)
    return f"{sec // 3600}:{sec % 3600 // 60:02d}:{sec % 60:02d}" if sec >= 3600 else f"{sec // 60}:{sec % 60:02d}"


def video_frame(path, t, max_side=640):
    """Один кадр в момент t (сек) как PIL.Image."""
    import av

    with av.open(path) as c:
        s = c.streams.video[0]
        s.thread_type = "AUTO"
        c.seek(int(t / s.time_base), stream=s, backward=True, any_frame=False)
        for fr in c.decode(s):
            if fr.time is None or fr.time >= t - 0.05:
                scale = min(1.0, max_side / max(fr.width, fr.height))
                w, h = max(2, int(fr.width * scale) // 2 * 2), max(2, int(fr.height * scale) // 2 * 2)
                return fr.reformat(width=w, height=h, format="rgb24").to_image()
    return None


def extract_video(path, cfg):
    import av

    with av.open(path) as c:
        if not c.streams.video:
            raise Skip("нет видеодорожки")
        s = c.streams.video[0]
        dur = c.duration / 1e6 if c.duration else float((s.duration or 0) * s.time_base)
    if dur <= 0:
        raise Skip("неизвестная длительность")
    step = max(cfg["video_frame_every_sec"], dur / cfg["video_max_frames"])
    times = [min(dur - 0.5, step / 2 + i * step) for i in range(cfg["video_max_frames"]) if step / 2 + i * step < dur]
    times = times or [dur / 2]
    items = []
    for t in times:
        try:
            fr = video_frame(path, t)
        except Exception:
            fr = None
        if fr is not None:
            items.append((t, _ts(t), "", fr))
    if not items:
        raise Skip("не удалось прочитать кадры")
    return items


def media_duration(path):
    """Длительность аудио/видео в секундах или None."""
    import av

    try:
        with av.open(path) as c:
            if c.duration:
                return c.duration / 1e6
            s = (c.streams.video or c.streams.audio)[0]
            return float(s.duration * s.time_base) if s.duration else None
    except Exception:
        return None


def extract_audio(path, cfg):
    import av

    seg = cfg["audio_segment_sec"] * 16000
    limit = seg * cfg["audio_max_segments"]
    bufs, total = [], 0
    with av.open(path) as c:
        if not c.streams.audio:
            raise Skip("нет аудио")
        rs = av.AudioResampler(format="flt", layout="mono", rate=16000)
        for fr in c.decode(c.streams.audio[0]):
            for r in rs.resample(fr):
                a = r.to_ndarray().reshape(-1)
                bufs.append(a)
                total += len(a)
            if total >= limit:
                break
    if not bufs:
        raise Skip("пустое аудио")
    a = np.concatenate(bufs)[:limit].astype(np.float32)
    items = []
    for i in range(0, len(a), seg):
        part = a[i:i + seg]
        if len(part) < 3 * 16000 and items:
            break
        items.append((i / 16000, _ts(i / 16000), "", part))
    return items


EXTRACTORS = {
    "text": extract_text,
    "code": extract_text,
    "image": extract_image,
    "video": extract_video,
    "audio": extract_audio,
}


def file_hash(path, kind):
    """Хеш содержимого для поиска точных копий. Аудио/видео и очень большие файлы — по размеру,
    началу и концу (читать гигабайты ради копий незачем)."""
    h = hashlib.sha1()
    size = os.path.getsize(path)
    with open(path, "rb") as f:
        if kind in ("audio", "video") or size > 64 * 2**20:
            h.update(str(size).encode())
            h.update(f.read(2**20))
            if size > 2**21:
                f.seek(-2**20, 2)
                h.update(f.read(2**20))
        else:
            for block in iter(lambda: f.read(2**20), b""):
                h.update(block)
    return f"{kind}:{h.hexdigest()}"


# Начало файла кода: «сгенерировано автоматически» — человек это не писал;
# шапка библиотеки (/*! … */, @license) — скорее всего чужой код, вложенный в проект
GEN_HEAD = re.compile(rb"auto-?generated|do not edit|@generated|generated by|this file (?:is|was) (?:auto-?)?generated",
                      re.I)
LIB_BANNER = re.compile(rb"^\s*/\*!|@license|@preserve", re.M)


def run_job(kind, path, cfg, prio=0):
    """-> (items, err, flag). flag: None; "skip" — пропустить правилом (status 3, причина в err);
    "defer" — отложить в последнюю очередь (prio 1). Верхнего уровня, чтобы запускаться в отдельных процессах."""
    mode = cfg.get("foreign_code", "skip")
    if kind == "code" and mode != "normal":
        try:
            with open(path, "rb") as f:
                head = f.read(2048)
        except OSError:
            head = b""
        if GEN_HEAD.search(head):
            if mode == "skip":
                return None, "сгенерированный файл", "skip"
            if prio == 0:
                return None, "сгенерированный файл", "defer"
        elif prio == 0 and LIB_BANNER.search(head):
            return None, "похоже на чужую библиотеку", "defer"
    try:
        return EXTRACTORS[kind](path, cfg), None, None
    except Skip as s:
        return None, str(s), None
    except Exception as ex:
        return None, f"{type(ex).__name__}: {ex}"[:200], None


def worker_init():
    """Рабочий процесс разбора PDF/документов: без шума предупреждений pypdf и в фоновом режиме
    (CPU idle, очень низкий приоритет диска и памяти) — игры и рабочий стол разбора не замечают."""
    import logging
    import warnings

    logging.disable(logging.WARNING)
    warnings.filterwarnings("ignore")
    try:
        import throttle

        throttle.background_process()
    except Exception:
        pass
